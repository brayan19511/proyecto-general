"""Lectura de los .xlsx de las cargas masivas (mismo formato que proyecto-05).

Primera fila = encabezados. Columnas reconocidas (sin distinguir mayúsculas):
- sku: obligatoria.
- on/off: active, is_active, on, on_off, on/off o estado. Valores: on/off,
  true/false, 1/0, si/no, activo/inactivo.

No se guarda el archivo: solo sus filas normalizadas y su SHA-256.
"""

from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.services.errors import InvalidDataError

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000

SKU_HEADERS = {"sku"}
ACTIVE_HEADERS = {"active", "is_active", "on", "on_off", "on/off", "estado"}
TRUE_VALUES = {"1", "true", "on", "yes", "si", "sí", "activo"}
FALSE_VALUES = {"0", "false", "off", "no", "inactivo"}


@dataclass
class ParsedFile:
    filename: str
    sha256: str
    rows: list[dict] = field(default_factory=list)  # [{"sku": str, "active": bool?}]
    errors: list[dict] = field(default_factory=list)  # [{"row", "message"}]


def read_upload(content: bytes, filename: str | None, *, active_column: str) -> ParsedFile:
    """active_column: "required", "optional" (sin columna = on) o "none" (se ignora)."""
    name = Path(filename or "carga.xlsx").name
    if Path(name).suffix.casefold() != ".xlsx":
        raise InvalidDataError("Solo se admiten archivos .xlsx.")
    if len(content) > MAX_FILE_BYTES:
        raise InvalidDataError("El archivo supera el límite de 10 MB.")
    if not content:
        raise InvalidDataError("El archivo está vacío.")

    parsed = ParsedFile(filename=name, sha256=sha256(content).hexdigest())
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, KeyError) as exc:
        raise InvalidDataError("El archivo no es un .xlsx válido.") from exc

    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [_header(value) for value in next(rows, None) or ()]
        sku_index = _find(headers, SKU_HEADERS)
        active_index = _find(headers, ACTIVE_HEADERS) if active_column != "none" else None
        if sku_index is None:
            parsed.errors.append({"row": 1, "message": "Falta la columna sku."})
        if active_column == "required" and active_index is None:
            parsed.errors.append({"row": 1, "message": "Falta la columna on/off (active, on_off o estado)."})
        if parsed.errors:
            return parsed

        seen: dict[str, int] = {}
        received = 0
        for number, values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in values):
                continue  # Filas vacías: se ignoran.
            received += 1
            if received > MAX_ROWS:
                parsed.errors.append({"row": number, "message": f"El archivo supera el límite de {MAX_ROWS} filas."})
                break

            sku = _sku(values[sku_index] if sku_index < len(values) else None)
            if not sku or len(sku) > 255:
                parsed.errors.append({"row": number, "message": "El SKU debe tener entre 1 y 255 caracteres."})
                continue
            key = sku.casefold()
            if key in seen:
                parsed.errors.append({"row": number, "message": f"SKU {sku} repetido (primera vez en la fila {seen[key]})."})
                continue
            seen[key] = number

            row: dict = {"sku": sku}
            if active_column != "none":
                if active_index is None:
                    row["active"] = True  # Carga completa sin columna on/off: lo listado queda on.
                else:
                    active = _active(values[active_index] if active_index < len(values) else None)
                    if active is None:
                        parsed.errors.append({"row": number, "message": f"Estado inválido para {sku}: usa on/off, 1/0 o true/false."})
                        continue
                    row["active"] = active
            parsed.rows.append(row)

        if received == 0:
            parsed.errors.append({"row": None, "message": "El archivo no tiene filas de datos."})
        return parsed
    finally:
        workbook.close()


def _header(value) -> str:
    return str(value or "").strip().casefold().replace(" ", "_")


def _find(headers: list[str], names: set[str]) -> int | None:
    return next((index for index, value in enumerate(headers) if value in names), None)


def _sku(value) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel guarda 123 como 123.0.
    text = str(value).strip()
    return text or None


def _active(value) -> bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        text = value.strip().casefold()
        if text in TRUE_VALUES:
            return True
        if text in FALSE_VALUES:
            return False
    return None
