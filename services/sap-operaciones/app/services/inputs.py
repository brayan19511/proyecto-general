"""Lectura y validación de las entradas de un trabajo (antes de crearlo).

- Documentos (acciones): texto con DocEntry separados por saltos de línea,
  comas, punto y coma o espacios (lo que sale al copiar una columna de Excel).
- Conciliaciones: .xlsx con las columnas de proyecto-05. Cada ReconNum es un
  ítem (un POST InternalReconciliations) y sus filas deben sumar 0.

Se juntan todos los errores (no solo el primero) para mostrarlos en la vista previa.
"""

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.services.errors import InvalidDataError

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ERRORS = 200  # Más errores que estos no aportan: se corta la lista.


@dataclass
class Parsed:
    items: list[tuple[str, dict | None]] = field(default_factory=list)  # (reference, payload)
    errors: list[dict] = field(default_factory=list)  # [{"row": n | None, "message": str}]
    duplicates: int = 0

    def error(self, row: int | None, message: str) -> None:
        if len(self.errors) < MAX_ERRORS:
            self.errors.append({"row": row, "message": message})


def parse_documents(text: str, max_items: int) -> Parsed:
    parsed = Parsed()
    seen: set[str] = set()
    for token in re.split(r"[\s,;]+", text.strip()):
        if not token:
            continue
        value = token[:-2] if token.endswith(".0") else token  # 123.0 copiado de Excel.
        if not value.isdigit() or int(value) <= 0:
            parsed.error(None, f"'{token[:30]}' no es un número de documento (DocEntry) válido.")
            continue
        value = str(int(value))  # Sin ceros a la izquierda: 00012 y 12 son el mismo documento.
        if value in seen:
            parsed.duplicates += 1
            continue
        seen.add(value)
        parsed.items.append((value, None))
    if not parsed.items and not parsed.errors:
        parsed.error(None, "No hay documentos.")
    if len(parsed.items) > max_items:
        parsed.error(None, f"Máximo {max_items} documentos por trabajo (hay {len(parsed.items)}).")
    return parsed


RECON_COLUMNS = ("ReconNum", "SrcObjAbs", "SrcObjTyp", "TransRowId", "TransId", "ShortName", "ReconcileAmount")


def parse_reconciliations(content: bytes, filename: str | None, recon_date: date, max_items: int) -> Parsed:
    """Excel de conciliación interna → un payload de InternalReconciliations por ReconNum."""
    if Path(filename or "").suffix.casefold() != ".xlsx":
        raise InvalidDataError("Solo se admiten archivos .xlsx.")
    if len(content) > MAX_FILE_BYTES:
        raise InvalidDataError("El archivo supera el límite de 20 MB.")
    if not content:
        raise InvalidDataError("El archivo está vacío.")
    try:
        # data_only=True: valor calculado de las celdas con fórmulas.
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, KeyError, ValueError) as exc:
        raise InvalidDataError("El archivo no es un .xlsx válido.") from exc

    parsed = Parsed()
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows, None) or ()
        columns = {str(value).strip(): index for index, value in enumerate(header) if value is not None and str(value).strip()}
        missing = [name for name in RECON_COLUMNS if name not in columns]
        if missing:
            parsed.error(1, "Faltan columnas obligatorias: " + ", ".join(missing) + ".")
            return parsed

        groups: dict[str, list[dict]] = defaultdict(list)
        totals: dict[str, Decimal] = defaultdict(Decimal)
        for number, values in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            cell = lambda name: values[columns[name]] if columns[name] < len(values) else None  # noqa: E731
            try:
                recon_num = _text(cell("ReconNum"), "ReconNum")
                amount = _decimal(cell("ReconcileAmount"))
                if amount == 0:
                    raise ValueError("ReconcileAmount no puede ser 0")
                line = {
                    "CashDiscount": None,
                    # SAP recibe montos absolutos; el lado lo define el signo del Excel.
                    "CreditOrDebit": "codDebit" if amount > 0 else "codCredit",
                    "ReconcileAmount": float(abs(amount)),
                    "Selected": "tYES",
                    "ShortName": _text(cell("ShortName"), "ShortName"),
                    "SrcObjAbs": _int(cell("SrcObjAbs"), "SrcObjAbs"),
                    "SrcObjTyp": _text(cell("SrcObjTyp"), "SrcObjTyp"),
                    "TransId": _int(cell("TransId"), "TransId"),
                    "TransRowId": _int(cell("TransRowId"), "TransRowId"),
                }
            except ValueError as exc:
                parsed.error(number, str(exc))
                continue
            groups[recon_num].append(line)
            totals[recon_num] += amount

        for recon_num, lines in groups.items():
            if abs(totals[recon_num]) > Decimal("0.01"):
                parsed.error(None, f"ReconNum {recon_num}: la suma de ReconcileAmount debe ser 0 (es {totals[recon_num]}).")
                continue
            parsed.items.append((
                recon_num[:100],
                {
                    "CardOrAccount": "coaCard",
                    "InternalReconciliationOpenTransRows": lines,
                    "ReconDate": recon_date.isoformat(),
                },
            ))
        if not groups and not parsed.errors:
            parsed.error(None, "El archivo no tiene filas para conciliar.")
        if len(parsed.items) > max_items:
            parsed.error(None, f"Máximo {max_items} conciliaciones por trabajo (hay {len(parsed.items)}).")
        return parsed
    finally:
        workbook.close()


def _text(value, name: str) -> str:
    if value is None or not str(value).strip():
        raise ValueError(f"{name} es obligatorio")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _int(value, name: str) -> int:
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, AttributeError):
        raise ValueError(f"{name} debe ser numérico") from None
    if value is None or number != number.to_integral_value():
        raise ValueError(f"{name} debe ser un número entero")
    return int(number)


def _decimal(value) -> Decimal:
    if value is None or not str(value).strip():
        raise ValueError("ReconcileAmount es obligatorio")
    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError("ReconcileAmount debe ser numérico") from None
